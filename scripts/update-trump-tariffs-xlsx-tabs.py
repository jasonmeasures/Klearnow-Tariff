#!/usr/bin/env python3
"""Add Sec 232 Pharma, Sec 201 QSP, and Sec 232 UAS summary tabs to the Trump tariffs workbook."""
from __future__ import annotations

import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "Trump_Tariffs_Summary_20260824 1.xlsx"


def main() -> int:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Font
    except ImportError:
        print("Install openpyxl: pip install openpyxl", file=sys.stderr)
        return 1

    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        print(f"Workbook not found: {xlsx}", file=sys.stderr)
        return 1

    wb = load_workbook(xlsx)
    bold = Font(bold=True)
    today = str(date.today())

    sheets = {
        "Sec 232 Pharma": [
            ["Primary HTS", "Sec 232", "Tariff Rate", "Source"],
            ["Ch.29 / Ch.30 (claim-gated)", "9903.04.60", "100% combined Col-1+232", "Proc. 11020 / CSMS #69395344"],
            ["JP / EU / KR / CH / LI patented", "9903.04.62", "15% combined", "CSMS #69395344"],
            ["UK patented (from 2026-07-31)", "9903.04.63", "0% additional", "CSMS #69415934"],
            ["Generic", "9903.04.67", "0%", "Proc. 11020"],
            ["", "", "", ""],
            ["Note", "Distinct from 301-FL pharma-use 9903.05.89 (Note 52(e)).", "", ""],
            ["Engine", "tariff-rules/data/s232_pharma.json", "", f"Tab added {today}"],
        ],
        "Sec 201 QSP": [
            ["Primary HTS", "Sec 201", "Tariff Rate", "Source"],
            ["6810.99.0020", "9903.45.30 / .31", "25% in-quota / 50% over-quota (Y1)", "U.S. note 41 / Proc. Jul 2026"],
            ["6810.99.0040", "9903.45.30 / .31", "TRQ through 2030-08-14", ""],
            ["7020.00.6000", "9903.45.30 / .31", "Stacks with 301-FL", ""],
            ["", "", "", ""],
            ["Engine", "tariff-rules/data/s201_qsp.json", "", f"Tab added {today}"],
        ],
        "Sec 232 UAS": [
            ["Primary HTS", "Sec 232", "Tariff Rate", "Effective", "Source"],
            ["8806.24/.29/.94/.99 (large UAS)", "9903.08.21", "100%", "2026-09-03", "Proc. 11055 / note 43"],
            ["8806.21–.23/.91–.93 (small UAS)", "9903.08.22", "25%", "2026-09-03", ""],
            ["8504.40.9580 / 8537.10.9170 (docking)", "9903.08.21", "100% (claim-gated)", "2026-09-03", ""],
            ["8807 parts", "9903.08.21 / .22", "claim-gated", "2026-09-03", ""],
            ["Partner caps", "9903.08.23 / .24", "10% / 15%", "2026-09-03", "combined mechanic TBC"],
            ["", "", "", "", ""],
            ["Engine", "tariff-rules/data/s232_uas.json", "", "", f"Tab added {today}"],
        ],
    }

    for name, rows in sheets.items():
        if name in wb.sheetnames:
            ws = wb[name]
            ws.delete_rows(1, ws.max_row)
        else:
            ws = wb.create_sheet(name)
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r_idx, column=c_idx, value=val)
                if r_idx == 1:
                    cell.font = bold

    toc = wb["ToC - Shortcuts"]
    start_row = toc.max_row + 1
    for i, title in enumerate(sheets):
        row = start_row + i
        toc.cell(row=row, column=1, value=f"{title}'!A1")
        prog = title.replace("Sec ", "")
        toc.cell(row=row, column=2, value=prog)
        toc.cell(row=row, column=3, value=f"Live pack summary — added {today}")

    wb.save(xlsx)
    print(f"Updated {xlsx.name} with tabs: {', '.join(sheets)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
